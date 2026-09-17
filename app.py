"""Interfaz ligera para Forecaster Futbol V10.4.2 Gatuno Adaptativo."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import html
import json
import os
from pathlib import Path
import traceback

import numpy as np
import pandas as pd
import streamlit as st

# La interfaz debe fallar rapido y conservar el resto de fuentes si un servidor
# externo no responde. El actualizador programado no define esta variable y
# puede usar los limites amplios de investigacion.
os.environ.setdefault("GATUNO_FAST_MODE", "1")

import bet_builder_v8_1_robust as base
import bet_forecaster_v10 as v10
import gatuno_audit as audit
import adaptive_monitor as adaptive
import model_monitor as quality_monitor


APP_VERSION = "V10.4.2-GATUNO-ADAPTATIVO"
DATA_DIR = Path("app_data_v10")
DATA_DIR.mkdir(exist_ok=True)
MATCH_FILE = DATA_DIR / "latest_matches.csv"
MARKET_FILE = DATA_DIR / "latest_markets.csv"
METRIC_FILE = DATA_DIR / "latest_validation.csv"
META_FILE = DATA_DIR / "meta.json"
RUN_STATE_FILE = DATA_DIR / "last_run_state.json"
METRICS_LOG_FILE = DATA_DIR / "model_metrics_log.csv"

st.set_page_config(
    page_title="V10.4.2 Gatuno Adaptativo",
    page_icon="🐾",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
    <style>
      .block-container {padding-top:.65rem; padding-bottom:3rem; max-width:1080px;}
      .hero {padding:1.05rem; border-radius:22px; border:1px solid rgba(128,128,128,.20);
             background:linear-gradient(135deg,rgba(20,184,166,.14),rgba(234,179,8,.11)); margin-bottom:.75rem;}
      .hero h1 {margin:0; font-size:clamp(1.45rem,6vw,2.15rem);}
      .hero p {margin:.42rem 0 0; opacity:.78;}
      .pill {display:inline-block; padding:.27rem .61rem; border-radius:999px; font-weight:850;
             font-size:.75rem; margin:.35rem .28rem 0 0;}
      .green {background:rgba(34,197,94,.17); color:#16a34a;}
      .amber {background:rgba(245,158,11,.18); color:#d97706;}
      .red {background:rgba(239,68,68,.16); color:#dc2626;}
      .blue {background:rgba(59,130,246,.15); color:#2563eb;}
      .gray {background:rgba(107,114,128,.15); color:#6b7280;}
      .audit {border:1px solid rgba(16,185,129,.22); border-radius:17px; padding:.78rem .9rem;
              background:rgba(16,185,129,.055); margin:.62rem 0;}
      .audit b {font-size:1.02rem;}
      .cat-panel {border:1px solid rgba(245,158,11,.25); border-radius:17px; padding:.82rem .92rem;
                  background:linear-gradient(135deg,rgba(245,158,11,.08),rgba(20,184,166,.07)); margin:.7rem 0;}
      .cat-row {padding:.55rem .65rem; border-radius:12px; margin:.35rem 0; background:rgba(128,128,128,.055);}
      .match-card {border:1px solid rgba(128,128,128,.20); border-radius:18px; padding:.95rem 1rem;
                   margin:.85rem 0; background:rgba(128,128,128,.035);}
      .match-title {font-weight:900; font-size:1.13rem; margin:.28rem 0 .15rem;}
      .sub {opacity:.68; font-size:.84rem;}
      .market-row {display:grid; grid-template-columns:1.42fr 1.15fr .72fr; gap:.42rem;
                   align-items:center; padding:.58rem .62rem; margin:.31rem 0; border-radius:11px;
                   background:rgba(128,128,128,.055); border-left:5px solid #6b7280;}
      .market-row.verde {border-left-color:#22c55e; background:rgba(34,197,94,.075);}
      .market-row.amarillo {border-left-color:#f59e0b; background:rgba(245,158,11,.075);}
      .market-row.rojo {border-left-color:#ef4444; background:rgba(239,68,68,.07);}
      .market-name {font-weight:780;}
      .pick {font-weight:900;}
      .signal {text-align:right; font-size:.77rem; font-weight:900;}
      .best {outline:2px solid rgba(34,197,94,.42);}
      .info-box {padding:.75rem .85rem; border-radius:14px; background:rgba(59,130,246,.08);
                 border:1px solid rgba(59,130,246,.18); margin:.65rem 0;}
      .stButton button,.stDownloadButton button {min-height:45px; border-radius:13px; font-weight:800;}
      @media(max-width:700px){
        .market-row {grid-template-columns:1.1fr .9fr;}
        .signal {grid-column:1/-1; text-align:left;}
      }
    </style>
    """,
    unsafe_allow_html=True,
)


def esc(value) -> str:
    return html.escape(str(value if value is not None else ""))


def as_bool(value) -> bool:
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    return str(value).strip().lower() in {"1", "true", "si", "sí", "yes"}


def atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    frame.to_csv(temporary, index=False, encoding="utf-8-sig")
    os.replace(temporary, path)


def atomic_json(payload: dict, path: Path) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)


def load_meta() -> dict:
    if not META_FILE.exists():
        return {}
    try:
        return json.loads(META_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def filter_not_started(frame: pd.DataFrame) -> pd.DataFrame:
    if frame is None or frame.empty:
        return pd.DataFrame() if frame is None else frame
    result = frame.copy()
    if "KickoffUTC" in result:
        kickoff = pd.to_datetime(result["KickoffUTC"], utc=True, errors="coerce")
        result = result[kickoff.notna() & (kickoff > pd.Timestamp.now(tz="UTC"))]
    return result.reset_index(drop=True)


def load_cached() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    matches = pd.read_csv(MATCH_FILE)
    markets = pd.read_csv(MARKET_FILE)
    metrics = pd.read_csv(METRIC_FILE)
    for frame in (matches, markets):
        if "Fecha" in frame:
            frame["Fecha"] = pd.to_datetime(frame["Fecha"], errors="coerce")
    return filter_not_started(matches), filter_not_started(markets), metrics


def refresh_everything(
    first_run: bool = False,
    progress=None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict, pd.DataFrame, dict]:
    # En la primera pantalla no existen pronosticos anteriores que cerrar.
    # Evitar esa consulta reduce latencia y elimina un punto de fallo remoto.
    if first_run:
        history = audit.load_history()
        closure = {"cerrados": 0, "pendientes": int(len(history)), "errores": []}
    else:
        if callable(progress):
            progress("Cerrando resultados anteriores que ya finalizaron…")
        history, closure = audit.resolve_pending()
    matches, markets, metrics, start, end = v10.run_v10(
        fast_mode=True,
        progress=progress,
    )
    matches = filter_not_started(matches)
    markets = filter_not_started(markets)
    markets = audit.adaptive_safety_gate(markets, history)
    markets = adaptive.apply_policy(markets, adaptive.load_policy())
    history, recording = audit.record_predictions(markets)
    adaptive.update_proposals(history)

    atomic_csv(matches, MATCH_FILE)
    atomic_csv(markets, MARKET_FILE)
    atomic_csv(metrics, METRIC_FILE)
    try:
        quality_monitor.registrar_metricas(metrics, METRICS_LOG_FILE)
    except Exception:
        pass
    meta = {
        "generated_at": datetime.now(base.TZ_PERU).isoformat(),
        "window_start": str(pd.Timestamp(start).date()),
        "window_end": str(pd.Timestamp(end).date()),
        "matches": int(len(matches)),
        "markets": int(len(markets)),
        "version": APP_VERSION,
        "audit_closed": int(closure.get("cerrados", 0)),
        "audit_inserted": int(recording.get("insertados", 0)),
    }
    atomic_json(meta, META_FILE)
    return matches, markets, metrics, meta, history, closure


def reapply_adaptive_policy(markets: pd.DataFrame) -> pd.DataFrame:
    """Recalcula la capa adaptativa desde el color previo, sin promociones."""
    base_markets = markets.copy()
    if "SemaforoPreAdaptativo" in base_markets:
        base_markets["SemaforoFinal"] = base_markets["SemaforoPreAdaptativo"]
        base_markets["Semaforo"] = base_markets["SemaforoPreAdaptativo"]
    if "AjusteAdaptativo" in base_markets:
        base_markets["AjusteAdaptativo"] = "SIN AJUSTE"
    adjusted = adaptive.apply_policy(base_markets, adaptive.load_policy())
    atomic_csv(adjusted, MARKET_FILE)
    return adjusted


def display_signal(signal: str) -> tuple[str, str]:
    value = str(signal).upper()
    return {
        "VERDE": ("🟢", "BUENA EVIDENCIA"),
        "AMARILLO": ("🟡", "PRECAUCIÓN"),
        "ROJO": ("🔴", "NO RECOMENDADA"),
    }.get(value, ("⚪", "SIN CLASIFICAR"))


def market_slug(code: str) -> str:
    return {
        "RESULT_1X2": "Resultado",
        "H1_GOALS_OU15": "Goles1T",
        "H1_CORNERS_OU45": "Corners1T",
        "YELLOW_CARDS_OU45": "Tarjetas",
        "HOME_SCORE_OU05": "GolLocal",
        "AWAY_SCORE_OU05": "GolVisitante",
    }.get(str(code), str(code))


def simple_summary(markets: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict] = []
    group_keys = [c for c in ["EventID", "KickoffUTC", "Local", "Visitante"] if c in markets]
    if not group_keys:
        return pd.DataFrame()
    for _, group in markets.groupby(group_keys, dropna=False, sort=False):
        first = group.iloc[0]
        item = {
            "Fecha": first.get("Fecha", ""),
            "HoraPeru": first.get("HoraPeru", ""),
            "Competicion": first.get("Competicion", ""),
            "Local": first.get("Local", ""),
            "Visitante": first.get("Visitante", ""),
        }
        for _, row in group.iterrows():
            slug = market_slug(row.get("MercadoCodigo", ""))
            item[f"{slug}_Pronostico"] = row.get("Pronostico", "")
            item[f"{slug}_Color"] = row.get("SemaforoFinal", row.get("Semaforo", ""))
        if "MejorOpcion" in group:
            best = group[group["MejorOpcion"].map(as_bool)]
        else:
            best = pd.DataFrame()
        item["MejorOpcionVerde"] = best.iloc[0].get("Pronostico", "") if not best.empty else "SIN SELECCION VERDE"
        rows.append(item)
    return pd.DataFrame(rows)


def excel_bytes(
    matches: pd.DataFrame,
    markets: pd.DataFrame,
    metrics: pd.DataFrame,
    history: pd.DataFrame,
) -> bytes:
    output = BytesIO()
    summary = simple_summary(markets)
    public_columns = [c for c in [
        "Fecha", "HoraPeru", "Competicion", "Local", "Visitante", "Mercado", "Linea",
        "Pronostico", "SemaforoFinal", "MejorOpcion", "EstadoAlineacion", "RiesgoRotacion",
        "AjusteAdaptativo",
    ] if c in markets]
    selections = markets[public_columns].copy()
    if "SemaforoFinal" not in selections and "Semaforo" in markets:
        selections["SemaforoFinal"] = markets["Semaforo"]

    audit_overview, audit_detail = audit.audit_summary(history)
    adaptive_diagnostics = adaptive.analyze_markets(history)
    adaptive_proposals = adaptive.load_proposals()
    audit_panel = pd.DataFrame([
        {"Indicador": "Estado", "Valor": audit_overview.get("state", "SIN_DATOS")},
        {"Indicador": "Pronosticos congelados", "Valor": audit_overview.get("total", 0)},
        {"Indicador": "Pendientes", "Valor": audit_overview.get("pending", 0)},
        {"Indicador": "Evaluados", "Valor": audit_overview.get("evaluated", 0)},
        {"Indicador": "Verdes evaluados", "Valor": audit_overview.get("green_evaluated", 0)},
        {"Indicador": "Tasa verde", "Valor": audit_overview.get("green_accuracy", np.nan)},
        {"Indicador": "Limite inferior 95%", "Valor": audit_overview.get("green_lcb95", np.nan)},
        {"Indicador": "Nota", "Valor": audit_overview.get("message", "")},
    ])

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="RESUMEN_COLORES", index=False)
        selections.to_excel(writer, sheet_name="SELECCIONES", index=False)
        audit_panel.to_excel(writer, sheet_name="AUDITORIA", index=False)
        if not audit_detail.empty:
            audit_detail.to_excel(writer, sheet_name="AUDITORIA_MERCADOS", index=False)
        if not adaptive_diagnostics.empty:
            adaptive_diagnostics.to_excel(writer, sheet_name="ALERTA_7_DIAS", index=False)
        if not adaptive_proposals.empty:
            adaptive_proposals.to_excel(writer, sheet_name="DECISIONES_AJUSTE", index=False)
        history.to_excel(writer, sheet_name="HISTORIAL", index=False)
        markets.to_excel(writer, sheet_name="TECNICO_MERCADOS", index=False)
        metrics.to_excel(writer, sheet_name="TECNICO_VALIDACION", index=False)

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = writer.book
        fills = {
            "VERDE": PatternFill("solid", fgColor="C6EFCE"),
            "AMARILLO": PatternFill("solid", fgColor="FFF2CC"),
            "ROJO": PatternFill("solid", fgColor="F4CCCC"),
        }
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column in range(1, sheet.max_column + 1):
                samples = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 80) + 1)]
                sheet.column_dimensions[get_column_letter(column)].width = min(42, max(11, max(map(len, samples), default=11) + 2))

        selection_sheet = workbook["SELECCIONES"]
        headers = {cell.value: cell.column for cell in selection_sheet[1]}
        signal_column = headers.get("SemaforoFinal")
        if signal_column:
            for row in range(2, selection_sheet.max_row + 1):
                status = str(selection_sheet.cell(row, signal_column).value or "")
                if status in fills:
                    for cell in selection_sheet[row]:
                        cell.fill = fills[status]

        summary_sheet = workbook["RESUMEN_COLORES"]
        summary_headers = {cell.value: cell.column for cell in summary_sheet[1]}
        for name, column in summary_headers.items():
            if str(name).endswith("_Color"):
                for row in range(2, summary_sheet.max_row + 1):
                    status = str(summary_sheet.cell(row, column).value or "")
                    if status in fills:
                        summary_sheet.cell(row, column).fill = fills[status]
                        if column > 1:
                            summary_sheet.cell(row, column - 1).fill = fills[status]

        for hidden_name in ("HISTORIAL", "TECNICO_MERCADOS", "TECNICO_VALIDACION"):
            workbook[hidden_name].sheet_state = "hidden"
    output.seek(0)
    return output.getvalue()


st.markdown(
    f"""
    <div class="hero">
      <h1>🐾😺 Forecaster Fútbol V10.4.2 Gatuno Adaptativo</h1>
      <p>Pronósticos individuales, cierre de jornada y ajustes trazables con tu confirmación.</p>
      <span class="pill green">🟢 BUENA EVIDENCIA</span>
      <span class="pill amber">🟡 PRECAUCIÓN</span>
      <span class="pill red">🔴 NO RECOMENDADA</span>
      <span class="pill blue">SIN CUOTAS INVENTADAS</span>
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="info-box"><b>Lectura simple:</b> cada partido muestra seis mercados separados.
    El sistema no arma combinadas, no usa hándicap asiático y no llama “verde” a un modelo que perdió
    contra su referencia temporal. Los cruces argentinos interliga se aprenden con datos comparables;
    no reciben un castigo fijo por nacionalidad. El Vigilante Gatuno separa resultado, goles, córners
    y tarjetas; si detecta deterioro repetido, te pide confirmar el ajuste o esperar siete días.</div>
    """,
    unsafe_allow_html=True,
)

meta = load_meta()
cache_ready = MATCH_FILE.exists() and MARKET_FILE.exists() and METRIC_FILE.exists() and meta.get("version") == APP_VERSION

if not cache_ready:
    last_run = {}
    if RUN_STATE_FILE.exists():
        try:
            last_run = json.loads(RUN_STATE_FILE.read_text(encoding="utf-8"))
        except Exception:
            last_run = {}
    if last_run.get("status") == "ERROR":
        st.error("La ejecución anterior falló y no se guardaron pronósticos.")
        st.code(str(last_run.get("error", "Error sin detalle")))
    elif last_run.get("status") == "RUNNING":
        st.warning(
            "La ejecución anterior fue interrumpida por el servidor antes de guardar la caché. "
            "V10.4.2 usa dos fuentes de calendario y un histórico autocontenido para evitar que vuelva a ocurrir."
        )
        if last_run.get("phase"):
            st.caption(f"Última fase registrada: {last_run['phase']}")
    else:
        st.warning("Primera ejecución rápida: genera una vez los datos. Después la pantalla abrirá desde la caché.")
    if st.button("🐾 Generar pronósticos V10.4.2", type="primary", use_container_width=True):
        try:
            atomic_json(
                {"status": "RUNNING", "phase": "Inicio", "updated_at": datetime.now(base.TZ_PERU).isoformat()},
                RUN_STATE_FILE,
            )
            with st.status("Construyendo V10.4.2 con respaldo autocontenido…", expanded=True) as status:
                def show_progress(message: str) -> None:
                    st.write(message)
                    atomic_json(
                        {"status": "RUNNING", "phase": message, "updated_at": datetime.now(base.TZ_PERU).isoformat()},
                        RUN_STATE_FILE,
                    )

                matches, markets, metrics, meta, history, closure = refresh_everything(
                    first_run=True,
                    progress=show_progress,
                )
                status.update(label="Pronósticos guardados correctamente", state="complete")
            atomic_json(
                {"status": "OK", "phase": "Completado", "updated_at": datetime.now(base.TZ_PERU).isoformat()},
                RUN_STATE_FILE,
            )
            st.rerun()
        except Exception as exc:
            atomic_json(
                {
                    "status": "ERROR",
                    "phase": "Primera ejecución",
                    "error": str(exc),
                    "updated_at": datetime.now(base.TZ_PERU).isoformat(),
                },
                RUN_STATE_FILE,
            )
            st.error("No fue posible completar la primera actualización.")
            st.code(str(exc))
            with st.expander("Detalle técnico"):
                st.code(traceback.format_exc())
    st.stop()

try:
    matches, markets, metrics = load_cached()
    history = audit.load_history()
except Exception as exc:
    st.error("La caché está dañada. Usa la actualización completa.")
    st.code(str(exc))
    st.stop()

audit_overview, audit_detail = audit.audit_summary(history)
accuracy = audit_overview.get("green_accuracy", np.nan)
accuracy_text = "SIN DATOS" if not np.isfinite(accuracy) else f"{accuracy:.1%}"
st.markdown(
    f"""
    <div class="audit"><b>📊 Auditoría real: {esc(audit_overview.get('state', 'SIN_DATOS'))}</b><br>
    Congelados: <b>{audit_overview.get('total', 0)}</b> · Pendientes: <b>{audit_overview.get('pending', 0)}</b> ·
    Evaluados: <b>{audit_overview.get('evaluated', 0)}</b> · Acierto verde: <b>{accuracy_text}</b><br>
    <span class="sub">{esc(audit_overview.get('message', ''))}</span></div>
    """,
    unsafe_allow_html=True,
)

# ---------------------------------------------------------------------------
# Vigilante Gatuno: alerta temprana y decisiones confirmadas por mercado.
# ---------------------------------------------------------------------------
try:
    adaptive_diagnostics, adaptive_proposals = adaptive.update_proposals(history)
except Exception as exc:
    adaptive_diagnostics, adaptive_proposals = pd.DataFrame(), pd.DataFrame()
    st.warning(f"El Vigilante Gatuno no pudo actualizarse: {exc}")

st.markdown(
    """
    <div class="cat-panel"><b>🐱🔎 Vigilante Gatuno de 7 días</b><br>
    <span class="sub">Busca errores repetidos por mercado. Una alerta no cambia nada por sí sola:
    tú decides aplicar una compuerta temporal o reunir siete días adicionales.</span></div>
    """,
    unsafe_allow_html=True,
)

if adaptive_diagnostics.empty:
    st.info("🐾 Aún no hay resultados verdes resueltos suficientes para evaluar mercados.")
else:
    state_view = {
        "ESTABLE": ("🟢", "Estable"),
        "OBSERVAR": ("🟡", "Observar"),
        "ALERTA": ("🟠", "Ajuste sugerido"),
        "CRITICA": ("🔴", "Pausa sugerida"),
        "SIN_MUESTRA": ("⚪", "Recolectando"),
    }
    for _, row in adaptive_diagnostics.iterrows():
        icon, label = state_view.get(str(row.get("Estado")), ("⚪", "Sin clasificar"))
        rate = row.get("TasaAcierto")
        expected = row.get("ProbMedia")
        rate_text = "sin tasa" if not np.isfinite(rate) else f"{rate:.0%} de acierto"
        expected_text = "" if not np.isfinite(expected) else f" · esperado {expected:.0%}"
        st.markdown(
            f"<div class='cat-row'>{icon} <b>{esc(row.get('Mercado',''))}</b> — {label} · "
            f"{int(row.get('N', 0))} verdes · {rate_text}{expected_text}</div>",
            unsafe_allow_html=True,
        )

pending_proposals = (
    adaptive_proposals[adaptive_proposals["Estado"].astype(str).eq("PENDIENTE")]
    if not adaptive_proposals.empty else pd.DataFrame()
)
if not pending_proposals.empty:
    st.error("🙀 Detecté un deterioro consistente. Revisa cada mercado antes de decidir.")
    for _, proposal in pending_proposals.iterrows():
        proposal_id = str(proposal["ProposalID"])
        with st.container(border=True):
            st.markdown(f"**🐾 {proposal['Mercado']} · {proposal['Severidad']}**")
            st.write(
                f"{int(proposal['Aciertos'])}/{int(proposal['N'])} aciertos "
                f"({proposal['TasaAcierto']:.1%}) frente a {proposal['ProbMedia']:.1%} esperado. "
                f"Propuesta: {str(proposal['AccionPropuesta']).replace('_', ' ').lower()}."
            )
            st.caption(str(proposal["Motivo"]))
            apply_col, wait_col = st.columns(2)
            if apply_col.button("🐾 Aplicar ajuste seguro", key=f"apply_{proposal_id}", use_container_width=True):
                adaptive.decide_proposal(proposal_id, "APLICAR")
                markets = reapply_adaptive_policy(markets)
                audit.sync_future_signals(markets)
                st.success("Ajuste temporal aplicado. Quedó registrado y puede revertirse.")
                st.rerun()
            if wait_col.button("🔎 Observar 7 días", key=f"watch_{proposal_id}", use_container_width=True):
                adaptive.decide_proposal(proposal_id, "OBSERVAR_7_DIAS")
                st.info("El criterio no cambió. Volveré a evaluarlo después de siete días.")
                st.rerun()

active_policies = adaptive.active_policy_rows()
if not active_policies.empty:
    with st.expander("😺 Ajustes Gatunos activos y reversión"):
        for _, rule in active_policies.iterrows():
            st.write(f"**{rule['Mercado']}** · {rule['Modo']} · vence {rule['Vence']}")
            if st.button("↩️ Revertir este ajuste", key=f"revert_{rule['MercadoCodigo']}"):
                adaptive.revert_market(str(rule["MercadoCodigo"]))
                markets = reapply_adaptive_policy(markets)
                audit.sync_future_signals(markets)
                st.rerun()

with st.expander("📐 Detalle técnico del Vigilante Gatuno"):
    st.caption(
        "La alerta exige al menos 20 verdes en 4 días y combina brecha de calibración, "
        "Brier, repetición diaria y comparación con 28 días previos. No usa solo una racha."
    )
    if not adaptive_diagnostics.empty:
        view = adaptive_diagnostics.copy()
        for column in ("TasaAcierto", "ProbMedia", "Brecha", "TasaPrevia"):
            if column in view:
                view[column] = view[column].apply(lambda value: "—" if not np.isfinite(value) else f"{value:.1%}")
        st.dataframe(view, use_container_width=True, hide_index=True)

try:
    metrics_log = pd.read_csv(METRICS_LOG_FILE) if METRICS_LOG_FILE.exists() else pd.DataFrame()
    quality_warnings = quality_monitor.detectar_degradacion(metrics_log) if not metrics_log.empty else []
except Exception:
    quality_warnings = []
if quality_warnings:
    st.warning(
        "⚠️ También cayó la calidad fuera de muestra de uno o más submodelos:\n\n"
        + "\n".join(f"- **{item['Modelo']}**: {item['Motivo']}" for item in quality_warnings)
    )

left, right = st.columns(2)
with left:
    if st.button("🔄 Actualizar y auditar resultados", use_container_width=True):
        try:
            with st.status("Actualizando V10.4.2…", expanded=True):
                matches, markets, metrics, meta, history, closure = refresh_everything(
                    first_run=False,
                    progress=st.write,
                )
                st.write(f"Pronósticos cerrados ahora: {closure.get('cerrados', 0)}")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))
            with st.expander("Detalle técnico"):
                st.code(traceback.format_exc())
with right:
    st.download_button(
        "📊 Descargar Excel claro",
        data=excel_bytes(matches, markets, metrics, history),
        file_name="V10_4_2_GATUNO_PRONOSTICOS_Y_AUDITORIA.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

with st.expander("💾 Respaldo de auditoría (opcional)"):
    st.caption(
        "Streamlit puede borrar archivos locales al reiniciar. Guarda este CSV y restáuralo después de un redespliegue para no perder la muestra prospectiva."
    )
    if not history.empty:
        st.download_button(
            "Descargar historial inmutable",
            data=history.to_csv(index=False, encoding="utf-8-sig"),
            file_name="historial_pronosticos_v1042.csv",
            mime="text/csv",
            use_container_width=True,
        )
    uploaded_history = st.file_uploader("Restaurar respaldo V10.3/V10.4", type=["csv"])
    if uploaded_history is not None and st.button("Restaurar sin sobrescribir registros existentes"):
        try:
            restored = pd.read_csv(uploaded_history, dtype={"PredictionID": "string", "EventID": "string"})
            required = {"PredictionID", "EstadoResultado", "EmitidoEnUTC", "MercadoCodigo"}
            if not required.issubset(restored.columns):
                raise ValueError("El archivo no tiene el esquema de historial Gatuno compatible.")
            allowed = {"PENDIENTE", "ACERTADO", "FALLADO", "NO_EVALUABLE"}
            if not set(restored["EstadoResultado"].astype(str)).issubset(allowed):
                raise ValueError("El respaldo contiene estados no válidos.")
            combined = pd.concat([history, restored], ignore_index=True, sort=False)
            combined = combined.drop_duplicates(subset=["PredictionID"], keep="first")
            atomic_csv(combined, audit.HISTORY_FILE)
            st.success(f"Historial restaurado: {len(combined)} registros únicos.")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))

generated = meta.get("generated_at", "")
try:
    generated = datetime.fromisoformat(generated).strftime("%d/%m/%Y %H:%M PET")
except Exception:
    pass
st.caption(
    f"Datos generados: {generated or '—'} · Ventana: {meta.get('window_start','—')} → {meta.get('window_end','—')} · {APP_VERSION}"
)

if matches.empty:
    st.warning("No quedan partidos futuros sin iniciar en la ventana actual. Actualiza el calendario.")
    st.stop()

f1, f2 = st.columns(2)
with f1:
    dates = sorted(pd.to_datetime(matches["Fecha"], errors="coerce").dropna().dt.date.unique())
    chosen_dates = st.multiselect("Fecha", dates, default=dates)
with f2:
    competitions = sorted(matches["Competicion"].dropna().astype(str).unique())
    chosen_competitions = st.multiselect("Competición", competitions, default=competitions)

visible_matches = matches[
    pd.to_datetime(matches["Fecha"], errors="coerce").dt.date.isin(chosen_dates)
    & matches["Competicion"].astype(str).isin(chosen_competitions)
]
visible_markets = markets[
    pd.to_datetime(markets["Fecha"], errors="coerce").dt.date.isin(chosen_dates)
    & markets["Competicion"].astype(str).isin(chosen_competitions)
]

green_count = int(visible_markets.get("SemaforoFinal", visible_markets.get("Semaforo", pd.Series(dtype=str))).eq("VERDE").sum())
best_count = int(visible_markets.get("MejorOpcion", pd.Series(False, index=visible_markets.index)).map(as_bool).sum())
m1, m2, m3 = st.columns(3)
m1.metric("Partidos", len(visible_matches))
m2.metric("Verdes", green_count)
m3.metric("Mejores opciones", best_count)

best_mask = visible_markets.get("MejorOpcion", pd.Series(False, index=visible_markets.index)).map(as_bool)
best = visible_markets[best_mask]
if not best.empty:
    st.subheader("🐱 Una mejor opción verde por partido")
    for _, row in best.iterrows():
        st.success(
            f"{row.get('Local','')} vs {row.get('Visitante','')}: "
            f"{row.get('Mercado','')} — {row.get('Pronostico','')}"
        )

st.subheader("Partidos y seis mercados")
for _, match in visible_matches.iterrows():
    event_id = str(match.get("EventID", ""))
    if event_id and event_id.lower() != "nan" and "EventID" in visible_markets:
        subset = visible_markets[visible_markets["EventID"].astype(str) == event_id]
    else:
        subset = visible_markets[
            (visible_markets["Local"].astype(str) == str(match.get("Local", "")))
            & (visible_markets["Visitante"].astype(str) == str(match.get("Visitante", "")))
            & (visible_markets["HoraPeru"].astype(str) == str(match.get("HoraPeru", "")))
        ]
    date_value = pd.Timestamp(match["Fecha"]).strftime("%d/%m/%Y") if pd.notna(match.get("Fecha")) else "—"
    arg_badge = '<span class="pill amber">Cruce argentino interliga</span>' if as_bool(match.get("CruceArgentinoInterliga", False)) else ""
    st.markdown(
        f"""
        <div class="match-card">
          <span class="pill blue">{esc(match.get('Competicion',''))}</span>
          <span class="pill gray">Rotación: {esc(match.get('RiesgoRotacion','—'))}</span>
          {arg_badge}
          <div class="match-title">{esc(match.get('Local',''))} vs {esc(match.get('Visitante',''))}</div>
          <div class="sub">{date_value} · {esc(match.get('HoraPeru',''))} PET · Alineación: {esc(match.get('EstadoAlineacion','—'))}</div>
        """,
        unsafe_allow_html=True,
    )
    for _, row in subset.iterrows():
        signal = str(row.get("SemaforoFinal", row.get("Semaforo", "ROJO"))).upper()
        emoji, label = display_signal(signal)
        best_class = " best" if as_bool(row.get("MejorOpcion", False)) else ""
        best_label = " · ⭐ MEJOR DEL PARTIDO" if as_bool(row.get("MejorOpcion", False)) else ""
        adaptive_note = str(row.get("AjusteAdaptativo", "") or "")
        adaptive_label = (
            f" · {esc(adaptive_note)}"
            if adaptive_note and adaptive_note not in {"SIN AJUSTE", "nan"}
            else ""
        )
        st.markdown(
            f"""
            <div class="market-row {signal.lower()}{best_class}">
              <div class="market-name">{esc(row.get('Mercado',''))} <small>({esc(row.get('Linea',''))})</small></div>
              <div class="pick">{esc(row.get('Pronostico',''))}</div>
              <div class="signal">{emoji} {label}{best_label}{adaptive_label}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    note = str(match.get("NotaContexto", "") or "")
    if note and note.lower() != "nan":
        st.caption("Contexto: " + note)
    st.markdown("</div>", unsafe_allow_html=True)

with st.expander("Cómo interpretar los colores"):
    st.markdown(
        """
        - **🟢 Verde:** superó referencia fuera de muestra, soporte, límite conservador y fiabilidad. No significa certeza.
        - **🟡 Amarillo:** señal intermedia; espera alineación y revisa el contexto.
        - **🔴 Rojo:** no apostar según el modelo: evidencia débil, cobertura insuficiente o validación fallida.
        - **⭐ Mejor del partido:** máximo una selección verde; nunca implica hacer una combinada.
        - La rentabilidad solo puede evaluarse con cuotas reales guardadas antes del partido.
        """
    )

st.caption("Sistema experimental. No garantiza aciertos ni beneficios; no persigas pérdidas y apuesta solo dinero que puedas perder.")
